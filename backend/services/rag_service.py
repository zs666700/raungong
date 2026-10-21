"""RAG：解析 → 清洗 → 分块 → 嵌入 → Chroma → 检索（BACKEND.md 第 17 章）。

同步流水线，不做异步任务与进度轮询：等就等了。

三件事必须守住：
1. **SQLite 是权威源，Chroma 只是索引**——命中的 document_id 必须回库校验 ready。
2. **嵌入必须串行**——模型不是线程安全的。
3. **写 Chroma 之前先按 document_id 清一次**，保证重跑从干净状态开始。
"""

from __future__ import annotations

import asyncio
import csv
import io
import ipaddress
import json
import logging
import math
import re
import shutil
import socket
import threading
import time
from collections import Counter
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import httpx
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from ..config import settings
from ..errors import ApiError
from ..models import RagDocument, User, utcnow

logger = logging.getLogger("tpa.rag")

COLLECTION_NAME = "tpa_knowledge"
ALLOWED_EXTS = {"pdf", "docx", "md", "txt", "csv"}
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_URL_BYTES = 5 * 1024 * 1024
CHUNK_OVERLAP = 120
MIN_CHUNK = 100
SCORE_THRESHOLD = 0.30
MAX_DOC_CHUNKS = 2   # 同一文档最多留 2 块
SNIPPET_CHARS = 200
MAX_TEXT_CHARS = 2_000_000

BLOCKED_NETWORKS = [
    ipaddress.ip_network(net)
    for net in (
        "127.0.0.0/8", "10.0.0.0/8", "172.16.0.0/12", "192.168.0.0/16",
        "169.254.0.0/16", "0.0.0.0/8", "::1/128", "fc00::/7", "fe80::/10",
    )
]

ZERO_WIDTH_RE = re.compile(r"[\u200b-\u200f\ufeff]")
MULTI_SPACE_RE = re.compile(r"[ \t\u3000]+")
MULTI_NEWLINE_RE = re.compile(r"\n{3,}")

# 嵌入
LOCAL_BATCH_SIZE = 16          # 本地模型：批量 16（17.5）
API_BATCH_SIZE = 10            # 线上 /embeddings：多数服务单请求上限 10–32 条，10 最稳
EMBED_TIMEOUT = 30.0           # 单次 HTTP 请求超时（秒）
RETRYABLE_STATUS = {429, 500, 502, 503, 504}

_embed_thread_lock = threading.Lock()


def _normalize(vector: list[float]) -> list[float]:
    """L2 归一化：cosine 距离下与归一化前的排序完全等价，但两个后端结果一致。

    本地模型本来就用 `normalize_embeddings=True`，线上接口不保证，所以统一在这里做。
    """
    norm = math.sqrt(sum(x * x for x in vector))
    if norm <= 0:
        return vector
    return [x / norm for x in vector]


# ---------------------------------------------------------------------------
# 校验 / 清洗 / 分块（纯函数，方便单测）
# ---------------------------------------------------------------------------


def assert_enabled() -> None:
    if not settings.rag_enabled:
        raise ApiError("RAG_DISABLED", "知识库未启用", 503)


def file_ext(filename: str) -> str:
    return Path(filename or "").suffix.lower().lstrip(".")


def safe_filename(filename: str) -> str:
    """剥离路径、过滤 `..` 和 `/`。"""
    name = Path(filename or "upload").name
    name = name.replace("\\", "_").replace("/", "_").replace("\x00", "")
    name = name.replace("..", "_").strip() or "upload"
    return name[:120]


def check_extension(filename: str) -> str:
    ext = file_ext(filename)
    if ext not in ALLOWED_EXTS:
        raise ApiError(
            "RAG_FILE_TYPE",
            f"只支持 {', '.join(sorted(ALLOWED_EXTS))} 格式",
            400,
        )
    return ext


def check_magic(ext: str, head: bytes, full_text_decodable: bool) -> None:
    """魔数不匹配就是改后缀骗人，直接拒。"""
    if ext == "pdf" and not head.startswith(b"%PDF-"):
        raise ApiError("RAG_FILE_TYPE", "文件内容不是合法的 PDF", 400)
    if ext == "docx" and not head.startswith(b"PK\x03\x04"):
        raise ApiError("RAG_FILE_TYPE", "文件内容不是合法的 DOCX", 400)
    if ext in {"md", "txt", "csv"} and not full_text_decodable:
        raise ApiError("RAG_FILE_TYPE", "文本文件无法按 UTF-8/GBK 解码", 400)


def clean_text(text: str) -> str:
    """17.3 清洗规则。"""
    if not text:
        return ""
    text = ZERO_WIDTH_RE.sub("", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [MULTI_SPACE_RE.sub(" ", line).strip() for line in text.split("\n")]
    text = "\n".join(lines)
    text = MULTI_NEWLINE_RE.sub("\n\n", text)
    return text.strip()


def drop_repeated_lines(pages: list[str]) -> str:
    """去页眉页脚：同一行在 >50% 页面重复出现则整行删除。"""
    if len(pages) < 3:
        return "\n\n".join(pages)
    counter: Counter[str] = Counter()
    for page in pages:
        for line in {line.strip() for line in page.split("\n") if line.strip()}:
            counter[line] += 1
    threshold = len(pages) / 2
    repeated = {line for line, count in counter.items() if count > threshold and len(line) < 80}
    if not repeated:
        return "\n\n".join(pages)
    cleaned = []
    for page in pages:
        kept = [line for line in page.split("\n") if line.strip() not in repeated]
        cleaned.append("\n".join(kept))
    return "\n\n".join(cleaned)


def split_text(text: str, chunk_size: int | None = None) -> list[str]:
    """800 字符、重叠 120、中文优先按句号切、最小块 100。"""
    size = max(200, int(chunk_size or settings.rag_chunk_size))
    text = (text or "").strip()
    if not text:
        return []

    chunks = _split_with_langchain(text, size)
    if chunks is None:
        chunks = _split_fallback(text, size)

    merged: list[str] = []
    for chunk in chunks:
        chunk = chunk.strip()
        if not chunk:
            continue
        if merged and len(chunk) < MIN_CHUNK:
            merged[-1] = f"{merged[-1]}\n{chunk}"  # 不足最小块与上一块合并
            continue
        merged.append(chunk)
    return merged


def _split_with_langchain(text: str, size: int) -> list[str] | None:
    splitter_cls = None
    try:
        from langchain_text_splitters import RecursiveCharacterTextSplitter

        splitter_cls = RecursiveCharacterTextSplitter
    except ImportError:
        try:
            from langchain.text_splitter import RecursiveCharacterTextSplitter

            splitter_cls = RecursiveCharacterTextSplitter
        except ImportError:
            return None
    splitter = splitter_cls(
        chunk_size=size,
        chunk_overlap=CHUNK_OVERLAP,
        length_function=len,
        separators=["\n\n", "\n", "。", "！", "？", "；", "，", " "],
    )
    return splitter.split_text(text)


def _split_fallback(text: str, size: int) -> list[str]:
    """没装 langchain 时的等价实现（按分隔符优先、带重叠）。"""
    separators = ["\n\n", "\n", "。", "！", "？", "；", "，", " "]
    pieces: list[str] = []
    cursor = 0
    while cursor < len(text):
        window = text[cursor : cursor + size]
        if len(window) < size:
            pieces.append(window)
            break
        cut = len(window)
        for sep in separators:
            pos = window.rfind(sep)
            if pos > size * 0.5:
                cut = pos + len(sep)
                break
        pieces.append(text[cursor : cursor + cut])
        cursor += max(1, cut - CHUNK_OVERLAP)
    return pieces


# ---------------------------------------------------------------------------
# 解析器（17.2）
# ---------------------------------------------------------------------------


def parse_pdf(data: bytes) -> tuple[str, list[dict[str, Any]]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # pragma: no cover
            logger.warning("PDF 某页解析失败：%s", exc)
            pages.append("")
    raw = drop_repeated_lines(pages)
    if len(raw.strip()) < 50:
        raise ApiError(
            "RAG_EMPTY_TEXT", "未能提取到文本，可能是扫描版 PDF，请先 OCR", 400
        )
    meta = [{"page": idx + 1} for idx, _ in enumerate(pages)]
    return raw, meta


def parse_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


def parse_text_file(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ApiError("RAG_EMPTY_TEXT", "文本文件无法按 UTF-8/GBK 解码", 400)


def parse_csv(data: bytes) -> str:
    text = parse_text_file(data)
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return ""
    header = rows[0]
    lines = []
    for row in rows[1:]:
        pairs = []
        for idx, value in enumerate(row):
            key = header[idx] if idx < len(header) else f"列{idx + 1}"
            pairs.append(f"{key}: {value}")
        lines.append("；".join(pairs))
    return "\n".join(lines)


def extract_plain_text(ext: str, data: bytes) -> str:
    if ext == "pdf":
        raw, _meta = parse_pdf(data)
        return raw
    if ext == "docx":
        return parse_docx(data)
    if ext == "csv":
        return parse_csv(data)
    return parse_text_file(data)


def extract_url_text(html: str) -> tuple[str, str | None]:
    """trafilatura 优先，失败退回 beautifulsoup4 取正文。"""
    title: str | None = None
    try:
        import trafilatura

        extracted = trafilatura.extract(
            html, include_comments=False, include_tables=True, url=None
        )
        try:
            from trafilatura.metadata import extract_metadata

            meta = extract_metadata(html)
            if meta and meta.title:
                title = meta.title.strip()
        except Exception:  # pragma: no cover
            pass
        if extracted and len(extracted.strip()) >= 50:
            return extracted, title
    except ImportError:
        logger.info("trafilatura 未安装，改用 beautifulsoup4")

    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    if title is None and soup.title and soup.title.string:
        title = soup.title.string.strip()
    for tag in soup(["script", "style", "nav", "header", "footer", "aside", "form"]):
        tag.decompose()
    body = soup.body or soup
    text = body.get_text("\n")
    return text, title


# ---------------------------------------------------------------------------
# SSRF 防护（13.2）
# ---------------------------------------------------------------------------


def assert_url_allowed(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https"):
        raise ApiError("RAG_URL_BLOCKED", "只允许 http/https", 400)
    host = parsed.hostname
    if not host:
        raise ApiError("RAG_URL_BLOCKED", "URL 缺少主机名", 400)
    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror as exc:
        raise ApiError("RAG_URL_BLOCKED", "无法解析该域名", 400) from exc
    for info in infos:
        raw_ip = info[4][0]
        try:
            ip = ipaddress.ip_address(raw_ip)
        except ValueError:
            continue
        if any(ip in network for network in BLOCKED_NETWORKS):
            raise ApiError("RAG_URL_BLOCKED", "该地址不允许抓取", 400)
    return url


# ---------------------------------------------------------------------------
# Chroma + 嵌入
# ---------------------------------------------------------------------------


class RagService:
    """整个进程用一个实例（`app.state.rag`）。

    嵌入有两种后端，由 `TPA_RAG_EMBEDDING_PROVIDER` 决定：

    * `api`   —— POST {base_url}/embeddings（OpenAI 兼容），只要 httpx；
    * `local` —— 本地 sentence-transformers 跑 BAAI/bge-m3（需 torch + 2.3 GB 权重）。

    两条路产出的向量维度必须与 `TPA_RAG_EMBEDDING_DIM` 一致（默认 1024）。
    """

    def __init__(self) -> None:
        self._collection: Any = None
        self._model: Any = None
        self._load_error: str | None = None
        self._asyncio_lock = asyncio.Lock()
        self._http: Any = None

    # ------------------------------------------------------------ 懒加载

    @property
    def enabled(self) -> bool:
        return settings.rag_enabled

    def _ensure_chroma(self) -> Any:
        if self._collection is not None:
            return self._collection
        try:
            import chromadb

            settings.chroma_dir.mkdir(parents=True, exist_ok=True)
            client = chromadb.PersistentClient(path=str(settings.chroma_dir))
            self._collection = client.get_or_create_collection(
                name=COLLECTION_NAME, metadata={"hnsw:space": "cosine"}
            )
        except Exception as exc:
            self._load_error = f"向量库不可用：{str(exc)[:200]}"
            logger.error("初始化 Chroma 失败：%s", exc)
            raise ApiError("RAG_DISABLED", self._load_error, 503) from exc
        return self._collection

    def _ensure_model(self) -> Any:
        """只在 provider=local 时用。"""
        if self._model is not None:
            return self._model
        try:
            from sentence_transformers import SentenceTransformer

            logger.info("加载本地嵌入模型 %s（首次可能很慢）", settings.rag_embedding_model)
            self._model = SentenceTransformer(settings.rag_embedding_model)
        except Exception as exc:
            self._load_error = f"本地嵌入模型不可用：{str(exc)[:200]}（可改用 TPA_RAG_EMBEDDING_PROVIDER=api）"
            logger.error("加载本地嵌入模型失败：%s", exc)
            raise ApiError("RAG_DISABLED", self._load_error, 503) from exc
        return self._model

    def _embed_local(self, texts: list[str]) -> list[list[float]]:
        """本地嵌入，串行（模型不是线程安全的）。"""
        model = self._ensure_model()
        with _embed_thread_lock:
            vectors = model.encode(
                texts,
                batch_size=LOCAL_BATCH_SIZE,
                normalize_embeddings=True,
                convert_to_numpy=True,
                show_progress_bar=False,
            )
        return [_normalize(list(map(float, vec))) for vec in vectors]

    async def embed(self, texts: list[str]) -> list[list[float]]:
        """统一的向量化入口：按 provider 分发，结果一律 L2 归一化。"""
        if not texts:
            return []
        async with self._asyncio_lock:
            if settings.embedding_is_api:
                return await self._embed_api(texts)
            return await asyncio.to_thread(self._embed_local, texts)

    # ------------------------------------------------------------ API 嵌入

    async def _client(self) -> Any:
        if self._http is None or self._http.is_closed:
            self._http = httpx.AsyncClient(timeout=httpx.Timeout(EMBED_TIMEOUT))
        return self._http

    async def aclose(self) -> None:
        if self._http is not None and not self._http.is_closed:
            await self._http.aclose()
            self._http = None

    async def _embed_api(self, texts: list[str]) -> list[list[float]]:
        """调 OpenAI 兼容的 /embeddings。分批 + 失败重试 1 次 + 维度校验。"""
        key = settings.rag_embedding_api_key.strip()
        if not key:
            raise ApiError(
                "RAG_DISABLED",
                "未配置 TPA_RAG_EMBEDDING_API_KEY，无法向量化（申请后填进 .env 并重启）",
                503,
            )

        url = settings.rag_embedding_base_url.rstrip("/") + "/embeddings"
        client = await self._client()
        vectors: list[list[float]] = []

        for start in range(0, len(texts), API_BATCH_SIZE):
            batch = texts[start : start + API_BATCH_SIZE]
            payload = {
                "model": settings.rag_embedding_model,
                "input": batch,
                "encoding_format": "float",
            }
            data = await self._post_embeddings(client, url, payload, key)
            if len(data) != len(batch):
                raise ApiError(
                    "RAG_DISABLED",
                    f"向量化服务返回 {len(data)} 条，与请求的 {len(batch)} 条不一致",
                    503,
                )
            # 有的服务会打乱顺序，按 index 摆回去
            ordered = sorted(
                enumerate(data),
                key=lambda pair: pair[1].get("index", pair[0])
                if isinstance(pair[1], dict)
                else pair[0],
            )
            for _pos, item in ordered:
                vector = item.get("embedding") if isinstance(item, dict) else None
                if not isinstance(vector, list) or not vector:
                    raise ApiError("RAG_DISABLED", "向量化服务返回了空向量", 503)
                if len(vector) != settings.rag_embedding_dim:
                    raise ApiError(
                        "RAG_DISABLED",
                        f"向量维度是 {len(vector)}，配置的 TPA_RAG_EMBEDDING_DIM 是 "
                        f"{settings.rag_embedding_dim}。换过模型的话请清空 data/chroma 后重传文档",
                        503,
                    )
                vectors.append(_normalize([float(x) for x in vector]))

        logger.info(
            "API 向量化完成：%s 条 / %s 批，模型=%s",
            len(vectors),
            (len(texts) + API_BATCH_SIZE - 1) // API_BATCH_SIZE,
            settings.rag_embedding_model,
        )
        return vectors

    async def _post_embeddings(
        self, client: Any, url: str, payload: dict[str, Any], key: str
    ) -> list[dict[str, Any]]:
        headers = {"Authorization": f"Bearer {key}", "Content-Type": "application/json"}
        last_error = ""

        for attempt in (1, 2):
            try:
                response = await client.post(url, json=payload, headers=headers)
            except httpx.HTTPError as exc:
                last_error = f"网络错误：{exc}"
                logger.warning("向量化请求失败（第 %s 次）：%s", attempt, exc)
                if attempt == 1:
                    await asyncio.sleep(0.8)
                    continue
                break

            if response.status_code in RETRYABLE_STATUS and attempt == 1:
                last_error = f"HTTP {response.status_code}：{response.text[:160]}"
                logger.warning("向量化服务 %s，重试一次", response.status_code)
                await asyncio.sleep(0.8)
                continue

            if response.status_code >= 400:
                detail = response.text[:200]
                if response.status_code in (401, 403):
                    hint = "（检查 TPA_RAG_EMBEDDING_API_KEY 是否正确、是否有该模型权限）"
                elif response.status_code == 404:
                    hint = "（检查 TPA_RAG_EMBEDDING_BASE_URL；一般要带 /v1，且模型名拼写正确）"
                elif response.status_code == 429:
                    hint = "（限流或余额不足）"
                else:
                    hint = ""
                raise ApiError(
                    "RAG_DISABLED",
                    f"向量化服务返回 HTTP {response.status_code}{hint}：{detail}",
                    503,
                )

            try:
                body = response.json()
            except ValueError as exc:
                raise ApiError("RAG_DISABLED", "向量化服务返回的不是 JSON", 503) from exc

            data = body.get("data")
            if not isinstance(data, list) or not data:
                raise ApiError(
                    "RAG_DISABLED", f"向量化服务返回格式异常：{str(body)[:200]}", 503
                )
            return data

        raise ApiError("RAG_DISABLED", f"向量化服务不可用：{last_error}", 503)

    # ------------------------------------------------------------ 入库

    async def _store_chunks(
        self,
        *,
        document_id: str,
        title: str,
        chunks: list[str],
        page_hint: list[dict[str, Any]] | None = None,
    ) -> None:
        collection = await asyncio.to_thread(self._ensure_chroma)
        vectors = await self.embed(chunks)
        ids = [f"{document_id}:{idx + 1}" for idx in range(len(chunks))]
        metadatas = []
        for idx in range(len(chunks)):
            meta: dict[str, Any] = {
                "document_id": document_id,
                "title": title[:200],
                "ordinal": idx + 1,
            }
            if page_hint and idx < len(page_hint) and page_hint[idx].get("page"):
                meta["page"] = int(page_hint[idx]["page"])
            metadatas.append(meta)
        await asyncio.to_thread(
            collection.add,
            ids=ids,
            documents=chunks,
            embeddings=vectors,
            metadatas=metadatas,
        )

    async def _purge_vectors(self, document_id: str) -> None:
        collection = await asyncio.to_thread(self._ensure_chroma)
        try:
            await asyncio.to_thread(
                collection.delete, where={"document_id": {"$eq": document_id}}
            )
        except Exception as exc:  # pragma: no cover
            logger.warning("清理向量失败 %s: %s", document_id, exc)

    async def ingest_file(
        self,
        db: AsyncSession,
        *,
        user: User,
        filename: str,
        title: str | None,
        data: bytes,
    ) -> dict[str, Any]:
        assert_enabled()
        ext = check_extension(filename)
        name = safe_filename(filename)
        doc = RagDocument(
            title=(title or Path(name).stem or name)[:200],
            source_type="file",
            filename=name,
            status="ready",
            chunk_count=0,
            created_by=user.id,
            created_at=utcnow(),
        )
        db.add(doc)
        await db.flush()

        doc_dir = settings.uploads_dir / doc.id
        doc_dir.mkdir(parents=True, exist_ok=True)
        (doc_dir / name).write_bytes(data)

        return await self._ingest_common(db, doc, ext, data)

    async def ingest_url(
        self, db: AsyncSession, *, user: User, url: str, title: str | None
    ) -> dict[str, Any]:
        assert_enabled()
        assert_url_allowed(url)
        html, page_title = await self._fetch_html(url)

        doc = RagDocument(
            title=(title or page_title or url)[:200],
            source_type="url",
            filename=url,
            status="ready",
            chunk_count=0,
            created_by=user.id,
            created_at=utcnow(),
        )
        db.add(doc)
        await db.flush()

        try:
            text, detected = await asyncio.to_thread(extract_url_text, html)
            detected = detected or page_title
        except Exception as exc:
            logger.warning("网页正文抽取失败 %s: %s", url, exc)
            text, detected = "", None
        if title is None and detected:
            doc.title = detected[:200]

        cleaned = clean_text(text)[:MAX_TEXT_CHARS]
        if len(cleaned) < 50:
            doc.status = "failed"
            doc.error = "网页没有可用的正文，可能被反爬或需要登录"
            await db.commit()
            return self._document_payload(doc, [])

        chunks = split_text(cleaned)
        return await self._finish_ingest(db, doc, chunks)

    async def _ingest_common(
        self, db: AsyncSession, doc: RagDocument, ext: str, data: bytes
    ) -> dict[str, Any]:
        try:
            text = await asyncio.to_thread(extract_plain_text, ext, data)
        except ApiError as exc:
            doc.status = "failed"
            doc.error = exc.message
            await db.commit()
            return self._document_payload(doc, [])

        cleaned = clean_text(text)
        if len(cleaned) > MAX_TEXT_CHARS:
            logger.warning("文档 %s 超长（%s 字符），截断到 200 万", doc.id, len(cleaned))
            cleaned = cleaned[:MAX_TEXT_CHARS]
        if len(cleaned) < 50:
            doc.status = "failed"
            doc.error = "没能提取到正文（清洗后不足 50 字符）"
            await db.commit()
            return self._document_payload(doc, [])

        chunks = split_text(cleaned)
        return await self._finish_ingest(db, doc, chunks)

    async def _finish_ingest(
        self, db: AsyncSession, doc: RagDocument, chunks: list[str]
    ) -> dict[str, Any]:
        # 写 Chroma 之前先按 document_id 清一次（保证重跑从干净状态开始）
        await self._purge_vectors(doc.id)
        try:
            await self._store_chunks(document_id=doc.id, title=doc.title, chunks=chunks)
        except ApiError as exc:
            doc.status = "failed"
            doc.error = exc.message
            await db.commit()
            return self._document_payload(doc, chunks)
        except Exception as exc:
            logger.exception("写入向量失败 %s", doc.id)
            await self._purge_vectors(doc.id)
            doc.status = "failed"
            doc.error = f"写入向量失败：{str(exc)[:200]}"
            await db.commit()
            return self._document_payload(doc, chunks)

        doc.status = "ready"
        doc.chunk_count = len(chunks)
        doc.error = None
        await db.commit()
        return self._document_payload(doc, chunks)

    @staticmethod
    def _document_payload(doc: RagDocument, chunks: list[str]) -> dict[str, Any]:
        return {
            "id": doc.id,
            "title": doc.title,
            "source_type": doc.source_type,
            "status": doc.status,
            "chunk_count": doc.chunk_count,
            "chunk_preview": [c[:100] for c in chunks[:3]],
            "error": doc.error,
        }

    async def _fetch_html(self, url: str) -> tuple[str, str | None]:
        """不跟随重定向（防公网 302 到内网），边读边计数，上限 5 MB。"""
        async with httpx.AsyncClient(timeout=15.0, follow_redirects=False) as client:
            try:
                async with client.stream("GET", url) as response:
                    if response.status_code >= 400:
                        raise ApiError(
                            "RAG_URL_BLOCKED",
                            f"抓取失败：HTTP {response.status_code}",
                            400,
                        )
                    chunks: list[bytes] = []
                    total = 0
                    async for piece in response.aiter_bytes():
                        total += len(piece)
                        if total > MAX_URL_BYTES:
                            raise ApiError(
                                "RAG_URL_BLOCKED", "网页内容超过 5 MB，已放弃", 400
                            )
                        chunks.append(piece)
                    body = b"".join(chunks)
                    encoding = response.encoding or "utf-8"
            except httpx.TimeoutException as exc:
                raise ApiError("RAG_URL_BLOCKED", "抓取超时", 400) from exc
            except httpx.HTTPError as exc:
                raise ApiError("RAG_URL_BLOCKED", f"抓取失败：{exc}", 400) from exc

        try:
            html = body.decode(encoding, errors="replace")
        except LookupError:
            html = body.decode("utf-8", errors="replace")
        return html, None

    # ------------------------------------------------------------ 列表 / 删除

    async def list_documents(self, db: AsyncSession) -> list[RagDocument]:
        assert_enabled()
        stmt = select(RagDocument).order_by(RagDocument.created_at.desc())
        return list((await db.execute(stmt)).scalars().all())

    async def delete_document(self, db: AsyncSession, doc_id: str) -> None:
        """顺序：删向量 -> 删行 -> 删磁盘目录。中途失败只记 WARNING。"""
        assert_enabled()
        doc = await db.get(RagDocument, doc_id)
        if doc is None:
            raise ApiError("NOT_FOUND", "文档不存在", 404)

        await self._purge_vectors(doc_id)
        await db.execute(delete(RagDocument).where(RagDocument.id == doc_id))
        await db.commit()

        doc_dir = settings.uploads_dir / doc_id
        if doc_dir.exists():
            try:
                await asyncio.to_thread(shutil.rmtree, doc_dir, True)
            except Exception as exc:  # pragma: no cover
                logger.warning("删除上传目录失败 %s: %s", doc_dir, exc)

    # ------------------------------------------------------------ 检索

    async def search(
        self, db: AsyncSession, query: str, top_k: int | None = None
    ) -> list[dict[str, Any]]:
        """检索 -> 阈值过滤 -> 按文档去重 -> 截断 -> 回库校验。"""
        assert_enabled()
        query = (query or "").strip()
        if not query:
            return []
        k = max(1, min(20, int(top_k or settings.rag_top_k)))

        collection = await asyncio.to_thread(self._ensure_chroma)
        vectors = await self.embed([query])

        try:
            result = await asyncio.to_thread(
                collection.query,
                query_embeddings=vectors,
                n_results=max(k * 2, k),
                include=["documents", "metadatas", "distances"],
            )
        except Exception as exc:
            logger.warning("向量检索失败：%s", exc)
            return []

        documents = (result.get("documents") or [[]])[0]
        metadatas = (result.get("metadatas") or [[]])[0]
        distances = (result.get("distances") or [[]])[0]

        candidates: list[dict[str, Any]] = []
        for text, meta, distance in zip(documents, metadatas, distances):
            meta = meta or {}
            try:
                score = 1.0 - float(distance)  # cosine 距离 -> 相似度
            except (TypeError, ValueError):
                continue
            if score < SCORE_THRESHOLD:
                continue
            document_id = str(meta.get("document_id") or "")
            if not document_id:
                continue
            candidates.append(
                {
                    "document_id": document_id,
                    "title": str(meta.get("title") or ""),
                    "ordinal": int(meta.get("ordinal") or 0),
                    "score": round(score, 4),
                    "snippet": (text or "")[:SNIPPET_CHARS],
                }
            )

        candidates.sort(key=lambda c: c["score"], reverse=True)

        # 同一文档最多留 2 块
        per_doc: Counter[str] = Counter()
        deduped: list[dict[str, Any]] = []
        for item in candidates:
            if per_doc[item["document_id"]] >= MAX_DOC_CHUNKS:
                continue
            per_doc[item["document_id"]] += 1
            deduped.append(item)

        # 二次校验：document_id 必须存在且 ready（防已删文档被喂给模型）
        if deduped:
            ids = {item["document_id"] for item in deduped}
            rows = (
                await db.execute(
                    select(RagDocument.id).where(
                        RagDocument.id.in_(ids), RagDocument.status == "ready"
                    )
                )
            ).scalars().all()
            alive = set(rows)
            deduped = [item for item in deduped if item["document_id"] in alive]

        return deduped[:k]

    async def search_payload(
        self, db: AsyncSession, query: str, top_k: int | None = None
    ) -> dict[str, Any]:
        """给 Agent 原生工具用的裸结构（15.8：只有 ok + hits，不套工具信封）。"""
        started = time.perf_counter()
        try:
            hits = await self.search(db, query, top_k)
        except ApiError as exc:
            if exc.code == "RAG_DISABLED":
                return {"ok": False, "hits": [], "error": exc.message}
            raise
        return {
            "ok": True,
            "hits": hits,
            "elapsed_ms": int((time.perf_counter() - started) * 1000),
        }

    def status(self) -> dict[str, Any]:
        """给 `/readyz` 和排障用：能一眼看出向量化走的是哪条路、配好没有。"""
        return {
            "enabled": self.enabled,
            "provider": settings.rag_embedding_provider,
            "model": settings.rag_embedding_model,
            "dim": settings.rag_embedding_dim,
            "base_url": (
                settings.rag_embedding_base_url if settings.embedding_is_api else None
            ),
            "configured": settings.embedding_configured,
            "load_error": self._load_error,
            "collection": COLLECTION_NAME,
        }


rag_service = RagService()
